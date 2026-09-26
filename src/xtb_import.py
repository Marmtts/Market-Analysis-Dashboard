"""
xtb_import.py
---------------
Import raportu XTB (.xlsx, arkusze "Open Positions" / "Closed Positions")
do portfela dashboardu. Bezpieczny i idempotentny: każda zaimportowana
pozycja dostaje znacznik [XTB:<id_pozycji>] w notatce - ponowny import
tego samego (lub nowszego) pliku NIGDY nie tworzy duplikatów, a istniejące
dane wpisane ręcznie w dashboardzie NIE są w żaden sposób ruszane.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime

import openpyxl

logger = logging.getLogger("xtb_trend_watch.xtb_import")

# Sufiksy XTB -> sufiksy Yahoo Finance. Nie jest to lista wyczerpująca dla
# wszystkich giełd świata - dla nierozpoznanego sufiksu ticker zostaje BEZ
# ZMIAN i trafia do listy "do ręcznej weryfikacji" w wyniku importu (można
# go potem poprawić przyciskiem ✎ w portfelu, dokładnie jak każdą inną pozycję).
XTB_SUFFIX_TO_YAHOO = {
    "US": "", "PL": "WA", "DE": "DE", "UK": "L", "FR": "PA", "NL": "AS",
    "ES": "MC", "IT": "MI", "PT": "LS", "BE": "BR", "AT": "VI", "CH": "SW",
    "SE": "ST", "NO": "OL", "DK": "CO", "FI": "HE", "CA": "TO", "JP": "T",
    "HK": "HK", "AU": "AX",
}

# Zgrubne domyślne przypisanie waluty wg sufiksu XTB - tylko jako punkt
# startowy (dokładna waluta i tak zostanie ustalona/skorygowana przez
# yfinance przy pierwszym pełnym cyklu analizy tej spółki).
XTB_SUFFIX_TO_CURRENCY = {
    "US": "USD", "PL": "PLN", "DE": "EUR", "FR": "EUR", "NL": "EUR",
    "ES": "EUR", "IT": "EUR", "PT": "EUR", "BE": "EUR", "AT": "EUR",
    "FI": "EUR", "UK": "GBP", "CH": "CHF", "SE": "SEK", "NO": "NOK",
    "DK": "DKK", "CA": "CAD", "JP": "JPY", "HK": "HKD", "AU": "AUD",
}


def map_xtb_ticker(xtb_ticker: str) -> tuple[str, bool]:
    """Zwraca (ticker_yahoo, czy_pewne_mapowanie). Przy nierozpoznanym
    sufiksie zwraca oryginalny ticker bez zmian i False - do oznaczenia
    jako 'wymaga weryfikacji' w wyniku importu."""
    if "." not in xtb_ticker:
        return xtb_ticker, False
    base, suffix = xtb_ticker.rsplit(".", 1)
    if suffix not in XTB_SUFFIX_TO_YAHOO:
        return xtb_ticker, False
    mapped_suffix = XTB_SUFFIX_TO_YAHOO[suffix]
    return (base if not mapped_suffix else f"{base}.{mapped_suffix}"), True


def guess_currency(xtb_ticker: str) -> str:
    if "." in xtb_ticker:
        suffix = xtb_ticker.rsplit(".", 1)[1]
        return XTB_SUFFIX_TO_CURRENCY.get(suffix, "USD")
    return "USD"


def _find_header(ws, required_columns: list[str]) -> tuple[int, dict[str, int]] | tuple[None, None]:
    """Szuka wiersza nagłówka po nazwach kolumn (nie po numerze wiersza -
    odporne na zmiany liczby wierszy podsumowania nad tabelą w różnych
    wariantach raportu XTB)."""
    for row in ws.iter_rows(min_row=1, max_row=25):
        values = [str(c.value).strip() if c.value is not None else "" for c in row]
        if all(col in values for col in required_columns):
            col_index = {name: i for i, name in enumerate(values)}
            return row[0].row, col_index
    return None, None


def _to_date_str(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return None


def _position_id_key(position_id) -> str | None:
    """Normalizuje Position ID (openpyxl zwraca float, np. 2798298883.0) do
    stałego formatu string, żeby dopasowywać wiersze między arkuszami."""
    if position_id in (None, ""):
        return None
    return str(int(position_id)) if isinstance(position_id, float) else str(position_id)


def _implied_fx_rate(account_ccy_amount, shares, native_price) -> float | None:
    """Kurs wymiany UKRYTY w kwocie rozliczonej w walucie konta - broker (np.
    XTB) dolicza własną marżę do kursu rynkowego, więc to jedyny sposób na
    odtworzenie kursu, jaki REALNIE zastosowano, zamiast zgadywać go kursem
    rynkowym/NBP z dnia transakcji."""
    if account_ccy_amount is None or not shares or not native_price:
        return None
    native_total = shares * native_price
    if not native_total:
        return None
    return abs(account_ccy_amount) / native_total


def _parse_cash_operations_buy_totals(wb) -> dict[str, float]:
    """Sumuje kwoty 'Stock purchase' z arkusza Cash Operations per Position ID
    - jedyne źródło rzeczywistego kosztu w walucie konta dla OTWARTYCH pozycji
    (arkusz Open Positions go nie ma). XTB czasem dzieli jeden zakup na kilka
    operacji gotówkowych (osobne partie konwersji waluty) - suma odtwarza
    całkowitą kwotę niezależnie od tego, na ile wierszy została podzielona."""
    totals: dict[str, float] = {}
    if "Cash Operations" not in wb.sheetnames:
        return totals
    ws = wb["Cash Operations"]
    header_row, cols = _find_header(ws, ["Type", "Amount", "Position ID"])
    if not header_row:
        return totals
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[cols["Type"]] != "Stock purchase":
            continue
        key = _position_id_key(row[cols["Position ID"]])
        amount = row[cols["Amount"]]
        if key is None or amount is None:
            continue
        totals[key] = totals.get(key, 0.0) + float(amount)
    return totals


def _parse_dividends(wb) -> tuple[list[dict], list[str]]:
    """Wyciąga wypłaty dywidend z arkusza Cash Operations: każda dywidenda to
    wiersz Type='Dividend' (kwota brutto w walucie konta), czasem sparowany z
    wierszem Type='Withholding tax' (podatek u źródła potrącony automatycznie
    przez brokera, kwota ujemna) - arkusz nie łączy ich żadnym wspólnym ID
    operacji, więc parujemy po (Ticker, Position ID), bo te DWA wiersze tej
    samej wypłaty zawsze mają identyczne Position ID (sam Time bywa o ułamek
    sekundy przesunięty między nimi - zweryfikowane na realnym eksporcie).
    Przy braku Position ID (nietypowy raport) spada na (Ticker, Time)."""
    warnings: list[str] = []
    if "Cash Operations" not in wb.sheetnames:
        return [], warnings
    ws = wb["Cash Operations"]
    header_row, cols = _find_header(ws, ["Type", "Ticker", "Time", "Amount", "ID"])
    if not header_row:
        return [], warnings
    has_position_id = "Position ID" in cols

    raw_rows = list(ws.iter_rows(min_row=header_row + 1, values_only=True))

    def pairing_key(row) -> tuple | None:
        ticker = row[cols.get("Ticker")]
        if ticker is None:
            return None
        ticker = str(ticker).strip()
        if has_position_id:
            pos_id = _position_id_key(row[cols["Position ID"]])
            if pos_id:
                return (ticker, "pos", pos_id)
        time_ = row[cols["Time"]]
        return (ticker, "time", time_) if time_ is not None else None

    wht_by_key: dict[tuple, float] = {}
    for row in raw_rows:
        if row[cols["Type"]] != "Withholding tax":
            continue
        key = pairing_key(row)
        amount = row[cols["Amount"]]
        if key is None or amount is None:
            continue
        wht_by_key[key] = wht_by_key.get(key, 0.0) + float(amount)

    dividend_rows: list[dict] = []
    for row in raw_rows:
        if row[cols["Type"]] != "Dividend":
            continue
        xtb_ticker, amount = row[cols.get("Ticker")], row[cols["Amount"]]
        op_id = row[cols.get("ID")] if "ID" in cols else None
        if xtb_ticker is None or amount is None:
            continue
        xtb_ticker = str(xtb_ticker).strip()

        yahoo_ticker, confident = map_xtb_ticker(xtb_ticker)
        if not confident:
            warnings.append(f"Nierozpoznany sufiks giełdy dla dywidendy {xtb_ticker} - "
                             f"zaimportowano jako '{yahoo_ticker}', ZWERYFIKUJ ręcznie.")

        key = pairing_key(row)
        wht = wht_by_key.get(key) if key else None
        comment = row[cols["Comment"]] if "Comment" in cols else None
        dividend_rows.append({
            "xtb_cash_op_id": str(int(op_id)) if isinstance(op_id, float) else (str(op_id) if op_id is not None else None),
            "ticker": yahoo_ticker,
            "original_ticker": xtb_ticker,
            "currency": guess_currency(xtb_ticker),
            "pay_date": _to_date_str(row[cols["Time"]]) or "",
            "amount_gross": round(float(amount), 6),
            "withholding_tax": round(abs(wht), 6) if wht else None,
            "notes": str(comment) if comment else "",
        })
    return dividend_rows, warnings


def parse_xtb_report(file_bytes: bytes) -> dict:
    """Zwraca {"open": [...], "closed": [...], "dividends": [...],
    "warnings": [...]}. Każdy element to dict gotowy do wstawienia do bazy
    portfela / tabeli dywidend."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    warnings: list[str] = []
    open_rows: list[dict] = []
    closed_rows: list[dict] = []

    # Kwoty rzeczywiście obciążające konto (waluta konta, z marżą brokera
    # wliczoną) dla OTWARTYCH pozycji - arkusz Open Positions samych kwot
    # historycznego kosztu nie ma, tylko bieżącą wartość.
    cash_buy_totals = _parse_cash_operations_buy_totals(wb)

    # --- Open Positions: tylko wiersze POJEDYNCZYCH transakcji (Type=BUY,
    # z wypełnioną datą otwarcia) - wiersze zagregowane (bez daty) pomijamy.
    if "Open Positions" in wb.sheetnames:
        ws = wb["Open Positions"]
        header_row, cols = _find_header(ws, ["Ticker", "Volume", "Open price", "Open time (UTC)"])
        if header_row:
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if row[cols["Ticker"]] is None:
                    continue
                position_id = row[cols.get("Instrument/Position", -1)] if "Instrument/Position" in cols else None
                open_time = row[cols["Open time (UTC)"]]
                type_ = row[cols.get("Type", -1)] if "Type" in cols else None
                if type_ != "BUY" or not open_time:
                    continue  # wiersz zagregowany albo krótka pozycja - pomijamy

                xtb_ticker = str(row[cols["Ticker"]]).strip()
                volume = row[cols["Volume"]]
                open_price = row[cols["Open price"]]
                if not volume or not open_price:
                    continue

                yahoo_ticker, confident = map_xtb_ticker(xtb_ticker)
                if not confident:
                    warnings.append(f"Nierozpoznany sufiks giełdy dla {xtb_ticker} - "
                                     f"zaimportowano jako '{yahoo_ticker}', ZWERYFIKUJ ręcznie.")

                id_key = _position_id_key(position_id)
                buy_fx_rate = _implied_fx_rate(cash_buy_totals.get(id_key), float(volume), float(open_price))

                open_rows.append({
                    "xtb_id": id_key,
                    "ticker": yahoo_ticker,
                    "original_ticker": xtb_ticker,
                    "shares": round(float(volume), 6),
                    "buy_price": round(float(open_price), 6),
                    "buy_date": _to_date_str(open_time) or datetime.now().strftime("%Y-%m-%d"),
                    "currency": guess_currency(xtb_ticker),
                    "buy_fx_rate": round(buy_fx_rate, 6) if buy_fx_rate else None,
                })

    # --- Closed Positions: każdy wiersz to już pojedyncza, zamknięta transakcja.
    # W przeciwieństwie do Open Positions, ten arkusz MA bezpośrednio kwoty w
    # walucie konta (Purchase Value / Sale Value) - nie trzeba sięgać do
    # Cash Operations.
    if "Closed Positions" in wb.sheetnames:
        ws = wb["Closed Positions"]
        header_row, cols = _find_header(ws, ["Ticker", "Open Price", "Close Price", "Position ID"])
        if header_row:
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if row[cols["Ticker"]] is None:
                    continue
                type_ = row[cols.get("Type", -1)] if "Type" in cols else "BUY"
                if type_ != "BUY":
                    warnings.append(f"Pominięto pozycję krótką (SELL) dla {row[cols['Ticker']]} - "
                                     f"nieobsługiwane w tym portfelu (tylko długie pozycje).")
                    continue

                xtb_ticker = str(row[cols["Ticker"]]).strip()
                volume = row[cols.get("Volume")]
                open_price = row[cols["Open Price"]]
                close_price = row[cols["Close Price"]]
                if not volume or not open_price or not close_price:
                    continue

                yahoo_ticker, confident = map_xtb_ticker(xtb_ticker)
                if not confident:
                    warnings.append(f"Nierozpoznany sufiks giełdy dla {xtb_ticker} - "
                                     f"zaimportowano jako '{yahoo_ticker}', ZWERYFIKUJ ręcznie.")

                purchase_value = row[cols["Purchase Value"]] if "Purchase Value" in cols else None
                sale_value = row[cols["Sale Value"]] if "Sale Value" in cols else None
                buy_fx_rate = _implied_fx_rate(purchase_value, float(volume), float(open_price))
                sell_fx_rate = _implied_fx_rate(sale_value, float(volume), float(close_price))

                position_id = row[cols["Position ID"]]
                closed_rows.append({
                    "xtb_id": _position_id_key(position_id),
                    "ticker": yahoo_ticker,
                    "original_ticker": xtb_ticker,
                    "shares": round(float(volume), 6),
                    "buy_price": round(float(open_price), 6),
                    "buy_date": _to_date_str(row[cols["Open Time (UTC)"]]) or "",
                    "sell_price": round(float(close_price), 6),
                    "sell_date": _to_date_str(row[cols["Close Time (UTC)"]]) or "",
                    "currency": guess_currency(xtb_ticker),
                    "buy_fx_rate": round(buy_fx_rate, 6) if buy_fx_rate else None,
                    "sell_fx_rate": round(sell_fx_rate, 6) if sell_fx_rate else None,
                })

    dividend_rows, dividend_warnings = _parse_dividends(wb)
    warnings.extend(dividend_warnings)

    return {"open": open_rows, "closed": closed_rows, "dividends": dividend_rows, "warnings": warnings}